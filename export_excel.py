"""
Writes the reconciliation result to a downloadable, formatted .xlsx:
  - Summary      : counts + value totals per category, plus run params
  - Exact Match  : side-by-side books/portal rows
  - Partial Match: side-by-side + a plain-English 'What to fix' note
  - Books Only   : present in our records, not in portal
  - Portal Only  : present in portal, not in our records

CANCEL / unresolved-GSTIN books rows are excluded earlier at extraction
time and reported on the extraction tab, not here.
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=10)
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14)
CELL_FONT = Font(name=FONT_NAME, size=10)
MISMATCH_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN_BORDER = Border(bottom=Side(style="thin", color="D9D9D9"))

STATUS_COLORS = {
    "Exact Match": "C6E0B4",
    "Partial Match": "FFE699",
    "Books Only": "F8CBAD",
    "Portal Only": "BDD7EE",
}

BOOKS_COLS = ["Doc Type", "GSTIN", "Party Name", "Doc No", "Doc Date",
              "Doc Value", "Taxable Value", "IGST", "CGST", "SGST", "CESS"]
PORTAL_COLS = BOOKS_COLS


def _style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autosize(ws, max_width=40):
    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=8)
        col_letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[col_letter].width = min(max(length + 2, 10), max_width)


def _write_side_by_side_sheet(wb, title, rows, status_label, include_mismatch_note=False):
    ws = wb.create_sheet(title)
    headers = (["Status"]
               + [f"Books - {c}" for c in BOOKS_COLS]
               + [f"Portal - {c}" for c in PORTAL_COLS])
    if include_mismatch_note:
        headers += ["Fields Requiring Correction", "What To Fix"]
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    ws.freeze_panes = "A2"

    r = 2
    for row in rows:
        b = row.books_doc.as_dict() if row.books_doc else {c: "" for c in BOOKS_COLS}
        p = row.portal_doc.as_dict() if row.portal_doc else {c: "" for c in PORTAL_COLS}
        vals = [status_label] + [b.get(c, "") for c in BOOKS_COLS] + [p.get(c, "") for c in PORTAL_COLS]
        if include_mismatch_note:
            field_list = ", ".join(m["field"] for m in row.mismatch_fields) if row.mismatch_fields else ""
            vals += [field_list, row.mismatch_note]
        ws.append(vals)
        for c in range(1, len(vals) + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = CELL_FONT
            cell.border = THIN_BORDER
        if include_mismatch_note and row.mismatch_fields:
            mismatched_field_names = {m["field"] for m in row.mismatch_fields}
            for c, colname in enumerate(BOOKS_COLS, start=2):
                key = {"Doc Value": "doc_value", "Taxable Value": "taxable_value", "IGST": "igst",
                       "CGST": "cgst", "SGST": "sgst", "CESS": "cess"}.get(colname)
                if key in mismatched_field_names:
                    ws.cell(row=r, column=c).fill = MISMATCH_FILL
                    ws.cell(row=r, column=c + len(BOOKS_COLS)).fill = MISMATCH_FILL
        r += 1

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(r - 1, 1)}"
    _autosize(ws)
    return ws


def _write_summary_sheet(wb, result, config, meta=None):
    ws = wb.create_sheet("Summary", 0)
    ws["A1"] = "GST Sales Reconciliation Summary"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:D1")

    r = 3
    if meta:
        for k, v in meta.items():
            ws.cell(row=r, column=1, value=k).font = Font(name=FONT_NAME, bold=True, size=10)
            ws.cell(row=r, column=2, value=v).font = CELL_FONT
            r += 1
        r += 1

    ws.cell(row=r, column=1, value="Match Tolerance (₹)").font = Font(name=FONT_NAME, bold=True, size=10)
    ws.cell(row=r, column=2, value=config.value_tolerance).font = CELL_FONT
    r += 1
    ws.cell(row=r, column=1, value="Date Tolerance (days, fallback matching)").font = Font(name=FONT_NAME, bold=True, size=10)
    ws.cell(row=r, column=2, value=config.date_tolerance_days).font = CELL_FONT
    r += 2

    header_row = r
    headers = ["Category", "Count", "Books Value (₹)", "Portal Value (₹)", "Description"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    r += 1

    descriptions = {
        "exact_match": ("Exact Match", "Books and portal agree within tolerance"),
        "partial_match": ("Partial Match (Correction Required)", "Same invoice found both sides but values differ"),
        "books_only": ("Present in Books, Not in Portal", "Sold/issued per our records, not yet reported/found on GSTR-1"),
        "portal_only": ("Present in Portal, Not in Books", "On GSTR-1 but missing from our records -- check for unrecorded sales or duplicate filing"),
    }

    for key in ["exact_match", "partial_match", "books_only", "portal_only"]:
        rows = result[key]
        label, desc = descriptions[key]
        books_val = sum(r_.books_doc.doc_value for r_ in rows if r_.books_doc)
        portal_val = sum(r_.portal_doc.doc_value for r_ in rows if r_.portal_doc)
        row_cells = [label, len(rows), round(books_val, 2), round(portal_val, 2), desc]
        status_key = label.split(" (")[0]
        for c, v in enumerate(row_cells, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = CELL_FONT
            if c == 1 and status_key in STATUS_COLORS:
                cell.fill = PatternFill("solid", fgColor=STATUS_COLORS[status_key])
        r += 1

    _autosize(ws, max_width=60)
    return ws


def export(result, config, filepath, meta=None):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    _write_summary_sheet(wb, result, config, meta)
    _write_side_by_side_sheet(wb, "Exact Match", result["exact_match"], "Exact Match")
    _write_side_by_side_sheet(wb, "Partial Match", result["partial_match"], "Partial Match", include_mismatch_note=True)
    _write_side_by_side_sheet(wb, "Books Only", result["books_only"], "Books Only")
    _write_side_by_side_sheet(wb, "Portal Only", result["portal_only"], "Portal Only")

    wb.save(filepath)
    return filepath
